from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import Workbook, load_workbook

from automations.recebimentos import (
    SourceRecord,
    analyze,
    payment_kind,
    reconcile,
    save_csv,
    save_excel,
)


class ReceiptsTests(TestCase):
    def test_payment_kind_normalizes_sources(self):
        self.assertEqual(payment_kind("Rede Visa 6"), ("Cartão de crédito", "Visa"))
        self.assertEqual(payment_kind("Mastercard", "débito"), ("Cartão de débito", "Mastercard"))
        self.assertEqual(payment_kind("Pix Hóspede"), ("Pix/Depósito", ""))

    def test_reconcile_matches_direct_folio_and_rede(self):
        day = date(2026, 7, 30)
        opera = [
            SourceRecord("OPERA", "10", day, "Dinheiro", "", Decimal("20.00"), folio="100"),
            SourceRecord("OPERA", "11", day, "Cartão de crédito", "Visa", Decimal("35.00"), last4="1234", folio="200"),
        ]
        rede = [SourceRecord("Rede", "999", day, "Cartão de crédito", "Visa", Decimal("35.00"), last4="1234")]
        cmflex = [
            SourceRecord("CMFlex", "10", day, "Dinheiro", "", Decimal("20.00")),
            SourceRecord("CMFlex", "200011", day, "A faturar", "", Decimal("0.00"), folio="200"),
        ]
        result = reconcile(opera, rede, cmflex)
        self.assertEqual(result.matched_count, 2)
        self.assertEqual(result.divergent_count, 0)

    def test_reconcile_reports_orphans_both_ways(self):
        day = date(2026, 7, 30)
        opera = [SourceRecord("OPERA", "1", day, "Dinheiro", "", Decimal("10.00"))]
        rede = [SourceRecord("Rede", "2", day, "Cartão de crédito", "Visa", Decimal("20.00"))]
        cmflex = [SourceRecord("CMFlex", "3", day, "Dinheiro", "", Decimal("30.00"))]
        result = reconcile(opera, rede, cmflex)
        self.assertEqual(result.divergent_count, 3)

    def test_reconcile_rejects_conflicting_card_last4(self):
        day = date(2026, 7, 30)
        opera = [
            SourceRecord(
                "OPERA", "1", day, "Cartão de crédito", "Visa",
                Decimal("100.00"), last4="1111",
            )
        ]
        rede = [
            SourceRecord(
                "Rede", "2", day, "Cartão de crédito", "Visa",
                Decimal("100.00"), last4="9999",
            )
        ]
        cmflex = [
            SourceRecord(
                "CMFlex", "1", day, "Cartão de crédito", "Visa",
                Decimal("100.00"),
            )
        ]

        result = reconcile(opera, rede, cmflex)

        self.assertEqual(result.matched_count, 0)
        self.assertEqual(result.divergent_count, 2)
        self.assertIn("ausente na Rede", result.rows[0].detail)

    def test_reconcile_rejects_unsettled_cmflex_document(self):
        day = date(2026, 7, 30)
        opera = [SourceRecord("OPERA", "1", day, "Dinheiro", "", Decimal("100.00"))]
        cmflex = [
            SourceRecord(
                "CMFlex", "1", day, "Dinheiro", "", Decimal("100.00"),
                settled=False,
            )
        ]

        result = reconcile(opera, [], cmflex)

        self.assertEqual(result.matched_count, 0)
        self.assertEqual(result.divergent_count, 1)
        self.assertIn("não baixado", result.rows[0].detail)

    def test_fiscal_document_can_cover_multiple_payments_from_same_folio(self):
        day = date(2026, 7, 30)
        opera = [
            SourceRecord("OPERA", "1", day, "Cartão de crédito", "Visa", Decimal("10.00"), folio="200"),
            SourceRecord("OPERA", "2", day, "Cartão de crédito", "Visa", Decimal("20.00"), folio="200"),
        ]
        rede = [
            SourceRecord("Rede", "10", day, "Cartão de crédito", "Visa", Decimal("10.00")),
            SourceRecord("Rede", "20", day, "Cartão de crédito", "Visa", Decimal("20.00")),
        ]
        cmflex = [SourceRecord("CMFlex", "200011", day, "A faturar", "", Decimal("0.00"), folio="200")]
        result = reconcile(opera, rede, cmflex)
        self.assertEqual(result.matched_count, 2)
        self.assertEqual(result.divergent_count, 0)

    def test_analyze_detects_all_three_sources_and_exports(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            opera = root / "opera.xml"
            opera.write_text(
                "<FINPAYMENTS><G_TRANSACTION><TRX_NO>10</TRX_NO><TRX_CODE>9086</TRX_CODE>"
                "<TRX_DESC>Dinheiro</TRX_DESC><FOLIO_NO>100</FOLIO_NO>"
                "<GUEST_ACCOUNT_CREDIT>-20</GUEST_ACCOUNT_CREDIT></G_TRANSACTION></FINPAYMENTS>",
                encoding="utf-8",
            )
            rede = root / "rede.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Extrato"])
            sheet.append(["data da venda", "status da venda", "valor da venda atualizado", "modalidade", "bandeira", "NSU/CV", "número da autorização (Auto)", "número do cartão"])
            sheet.append([day := date(2026, 7, 30), "negada", 0, "débito", "Visa", "50", "60", ""])
            workbook.save(rede)
            cmflex = root / "cmflex.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Numero", "Cliente", "PortadorForma", "TipoDeDocumento", "Valor", "DataEmissao", "StatusDoDocumento"])
            sheet.append(["10", "DINHEIRO", "CAIXA GERAL", "MOVIMENTO DE CAIXA", 20, day, "Baixado"])
            workbook.save(cmflex)

            result = analyze([rede, cmflex, opera])
            self.assertEqual(result.matched_count, 1)
            self.assertEqual(result.ignored_rede_count, 1)
            csv_path = root / "result.csv"
            excel_path = root / "result.xlsx"
            save_csv(result, csv_path)
            save_excel(result, excel_path)
            self.assertIn("ID OPERA", csv_path.read_text(encoding="utf-8-sig"))
            self.assertEqual(load_workbook(excel_path).sheetnames, ["Conferência", "Resumo"])
