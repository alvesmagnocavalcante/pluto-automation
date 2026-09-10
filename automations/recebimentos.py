from __future__ import annotations

import csv
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from defusedxml.ElementTree import parse as parse_xml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CENT = Decimal("0.01")


@dataclass(frozen=True)
class SourceRecord:
    source: str
    identifier: str
    business_date: date | None
    method: str
    brand: str
    amount: Decimal
    last4: str = ""
    folio: str = ""
    description: str = ""
    settled: bool = True


@dataclass(frozen=True)
class ReceiptRow:
    source: str
    date: date | None
    method: str
    brand: str
    amount: Decimal
    opera_id: str = ""
    rede_id: str = ""
    cmflex_id: str = ""
    status: str = ""
    detail: str = ""


@dataclass(frozen=True)
class ReconciliationResult:
    rows: tuple[ReceiptRow, ...]
    opera_count: int
    rede_count: int
    cmflex_count: int
    ignored_rede_count: int

    @property
    def matched_count(self) -> int:
        return sum(row.status == "OK" for row in self.rows)

    @property
    def divergent_count(self) -> int:
        return sum(row.status == "DIVERGENTE" for row in self.rows)


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in text if not unicodedata.combining(char)).casefold().strip()


def identifier(value: object) -> str:
    text = str(value or "").strip()
    return text[:-2] if text.endswith(".0") else text


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def decimal(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    text = str(value).strip().replace("R$", "").replace("\xa0", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text).quantize(CENT)
    except InvalidOperation as error:
        raise ValueError(f"Valor monetário inválido: {value}") from error


def card_last4(value: object) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-4:] if len(digits) >= 4 else ""


def payment_kind(description: object, modality: object = "") -> tuple[str, str]:
    text = normalize(f"{description} {modality}")
    brand = next((name for token, name in (("master", "Mastercard"), ("visa", "Visa"), ("elo", "Elo"), ("amex", "Amex")) if token in text), "")
    if "dinheiro" in text or "caixa geral" in text:
        return "Dinheiro", brand
    if "pix" in text or "deposito" in text or "extrato bancario" in text:
        return "Pix/Depósito", brand
    if "fatur" in text or "nota fiscal" in text:
        return "A faturar", brand
    if "debito" in text:
        return "Cartão de débito", brand
    if "credito" in text or brand:
        return "Cartão de crédito", brand
    return "Outro", brand


def worksheet_records(path: Path, header_row: int) -> list[dict[str, object]]:
    sheet = load_workbook(path, data_only=True, read_only=False).active
    headers = [str(sheet.cell(header_row, column).value or "").strip() for column in range(1, sheet.max_column + 1)]
    if not any(headers):
        raise ValueError(f"{path.name} não possui cabeçalho.")
    return [
        dict(zip(headers, (sheet.cell(row, column).value for column in range(1, sheet.max_column + 1))))
        for row in range(header_row + 1, sheet.max_row + 1)
        if any(sheet.cell(row, column).value not in (None, "") for column in range(1, sheet.max_column + 1))
    ]


def read_opera(path: Path) -> list[SourceRecord]:
    try:
        transactions = parse_xml(path).getroot().findall(".//G_TRANSACTION")
    except Exception as error:
        raise ValueError(f"XML do OPERA inválido: {error}") from error
    if not transactions:
        raise ValueError("O XML não contém transações do OPERA.")
    records = []
    for node in transactions:
        values = {child.tag: (child.text or "").strip() for child in node}
        method, brand = payment_kind(values.get("TRX_DESC"))
        records.append(
            SourceRecord(
                source="OPERA",
                identifier=identifier(values.get("TRX_NO")),
                business_date=parse_date(values.get("BUSINESS_DATE")),
                method=method,
                brand=brand,
                amount=abs(decimal(values.get("GUEST_ACCOUNT_CREDIT"))),
                last4=card_last4(values.get("CARD_NUMBER")),
                folio=identifier(values.get("FOLIO_NO")),
                description=values.get("TRX_DESC", ""),
            )
        )
    return records


def find_header(records: list[dict[str, object]], expected: str) -> str:
    expected_normalized = normalize(expected)
    for header in records[0]:
        if normalize(header) == expected_normalized:
            return header
    raise ValueError(f"Coluna obrigatória ausente: {expected}")


def read_rede(path: Path) -> list[SourceRecord]:
    records = worksheet_records(path, 2)
    if not records:
        raise ValueError("A planilha da Rede não possui vendas.")
    columns = {name: find_header(records, name) for name in (
        "data da venda", "status da venda", "valor da venda atualizado", "modalidade",
        "bandeira", "NSU/CV", "número da autorização (Auto)", "número do cartão",
    )}
    result = []
    for row in records:
        status = normalize(row[columns["status da venda"]])
        method, brand = payment_kind(row[columns["bandeira"]], row[columns["modalidade"]])
        card = row[columns["número do cartão"]]
        result.append(
            SourceRecord(
                source="Rede",
                identifier=identifier(row[columns["NSU/CV"]]) or identifier(row[columns["número da autorização (Auto)"]]),
                business_date=parse_date(row[columns["data da venda"]]),
                method=method,
                brand=brand,
                amount=abs(decimal(row[columns["valor da venda atualizado"]])),
                last4=card_last4(card),
                description=str(row[columns["status da venda"]] or ""),
                settled=status in {"aprovada", "aprovado", "pago", "paga"},
            )
        )
    return result


def read_cmflex(path: Path) -> list[SourceRecord]:
    records = worksheet_records(path, 1)
    if not records:
        raise ValueError("A planilha do CMFlex não possui lançamentos.")
    columns = {name: find_header(records, name) for name in (
        "Numero", "Cliente", "PortadorForma", "TipoDeDocumento", "Valor", "DataEmissao", "StatusDoDocumento",
    )}
    result = []
    for row in records:
        document = row[columns["TipoDeDocumento"]]
        method, brand = payment_kind(f"{document} {row[columns['PortadorForma']]}")
        number = identifier(row[columns["Numero"]])
        folio = number[:-3] if normalize(document).startswith("nota fiscal") and number.endswith("011") else ""
        result.append(
            SourceRecord(
                source="CMFlex",
                identifier=number,
                business_date=parse_date(row[columns["DataEmissao"]]),
                method=method,
                brand=brand,
                amount=decimal(row[columns["Valor"]]),
                folio=folio,
                description=f"{document} | {row[columns['Cliente']]}",
                settled=normalize(row[columns["StatusDoDocumento"]]) == "baixado",
            )
        )
    return result


def detect_sources(paths: Iterable[Path]) -> tuple[Path, Path, Path]:
    opera = rede = cmflex = None
    for path in map(Path, paths):
        if path.suffix.casefold() == ".xml":
            opera = path
            continue
        if path.suffix.casefold() != ".xlsx":
            raise ValueError(f"Formato não suportado: {path.name}")
        sheet = load_workbook(path, read_only=False, data_only=True).active
        first_rows = " ".join(str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 3) + 1) for column in range(1, min(sheet.max_column, 45) + 1))
        normalized = normalize(first_rows)
        if "nsu/cv" in normalized and "status da venda" in normalized:
            rede = path
        elif "portadorforma" in normalized and "tipodedocumento" in normalized:
            cmflex = path
    missing = [name for name, value in (("OPERA (.xml)", opera), ("Rede (.xlsx)", rede), ("CMFlex (.xlsx)", cmflex)) if value is None]
    if missing:
        raise ValueError("Arquivos não identificados: " + ", ".join(missing))
    return opera, rede, cmflex


def same_date(left: date | None, right: date | None) -> bool:
    return left is None or right is None or left == right


def match_cmflex(opera: SourceRecord, records: list[SourceRecord], used: set[int]) -> int | None:
    direct = [index for index, item in enumerate(records) if index not in used and item.identifier == opera.identifier]
    if direct:
        return direct[0]
    # Um documento fiscal do CMFlex pode consolidar mais de um recebimento do
    # mesmo fólio. Por isso, o vínculo por fólio é reutilizável; somente o
    # vínculo financeiro direto por número é estritamente um-para-um.
    by_folio = [index for index, item in enumerate(records) if item.folio and item.folio == opera.folio]
    if not by_folio:
        return None
    exact_unused = [index for index in by_folio if index not in used and records[index].amount == opera.amount]
    unused = [index for index in by_folio if index not in used]
    exact_value = [index for index in by_folio if records[index].amount == opera.amount]
    return (exact_unused or unused or exact_value or by_folio)[0]


def match_rede(opera: SourceRecord, records: list[SourceRecord], used: set[int]) -> int | None:
    candidates = [
        index for index, item in enumerate(records)
        if index not in used and item.settled and item.amount == opera.amount
        and item.method == opera.method and same_date(item.business_date, opera.business_date)
        and (not opera.brand or not item.brand or item.brand == opera.brand)
    ]
    if opera.last4:
        same_card = [index for index in candidates if records[index].last4 == opera.last4]
        without_card = [index for index in candidates if not records[index].last4]
        candidates = same_card or without_card
    return candidates[0] if candidates else None


def reconcile(opera: list[SourceRecord], rede: list[SourceRecord], cmflex: list[SourceRecord]) -> ReconciliationResult:
    used_rede: set[int] = set()
    used_cmflex: set[int] = set()
    rows: list[ReceiptRow] = []
    for source in opera:
        cm_index = match_cmflex(source, cmflex, used_cmflex)
        if cm_index is not None:
            used_cmflex.add(cm_index)
        needs_rede = source.method.startswith("Cartão") or source.method == "Pix/Depósito"
        rede_index = match_rede(source, rede, used_rede) if needs_rede else None
        if rede_index is not None:
            used_rede.add(rede_index)
        problems = []
        if cm_index is None:
            problems.append("ausente no CMFlex")
        else:
            cmflex_record = cmflex[cm_index]
            if not cmflex_record.settled:
                problems.append("documento não baixado no CMFlex")
            if cmflex_record.amount and abs(cmflex_record.amount) != source.amount:
                problems.append("valor divergente no CMFlex")
        if needs_rede and rede_index is None:
            problems.append("ausente na Rede")
        rows.append(
            ReceiptRow(
                source="OPERA",
                date=source.business_date,
                method=source.method,
                brand=source.brand,
                amount=source.amount,
                opera_id=source.identifier,
                rede_id=rede[rede_index].identifier if rede_index is not None else "",
                cmflex_id=cmflex[cm_index].identifier if cm_index is not None else "",
                status="DIVERGENTE" if problems else "OK",
                detail="; ".join(problems) if problems else "Lançamento conciliado",
            )
        )
    for index, source in enumerate(rede):
        if index in used_rede:
            continue
        ignored = not source.settled
        rows.append(ReceiptRow("Rede", source.business_date, source.method, source.brand, source.amount, rede_id=source.identifier, status="IGNORADO" if ignored else "DIVERGENTE", detail="Venda não liquidada" if ignored else "ausente no OPERA"))
    for index, source in enumerate(cmflex):
        if index not in used_cmflex:
            rows.append(ReceiptRow("CMFlex", source.business_date, source.method, source.brand, source.amount, cmflex_id=source.identifier, status="DIVERGENTE", detail="ausente no OPERA"))
    return ReconciliationResult(tuple(rows), len(opera), len(rede), len(cmflex), sum(not item.settled for item in rede))


def analyze(paths: list[Path]) -> ReconciliationResult:
    if len(paths) != 3:
        raise ValueError("Selecione exatamente os três arquivos: OPERA, Rede e CMFlex.")
    opera_path, rede_path, cmflex_path = detect_sources(paths)
    return reconcile(read_opera(opera_path), read_rede(rede_path), read_cmflex(cmflex_path))


EXPORT_HEADERS = ("Origem", "Data", "Meio", "Bandeira", "Valor", "ID OPERA", "ID Rede", "ID CMFlex", "Resultado", "Detalhe")


def export_values(row: ReceiptRow) -> tuple[object, ...]:
    return (row.source, row.date.strftime("%d/%m/%Y") if row.date else "", row.method, row.brand, row.amount, row.opera_id, row.rede_id, row.cmflex_id, row.status, row.detail)


def save_csv(result: ReconciliationResult, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(EXPORT_HEADERS)
        for row in result.rows:
            values = list(export_values(row))
            values[4] = str(values[4]).replace(".", ",")
            writer.writerow(values)


def save_excel(result: ReconciliationResult, output: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conferência"
    sheet.append(EXPORT_HEADERS)
    for row in result.rows:
        sheet.append(export_values(row))
    fill = PatternFill("solid", fgColor="263238")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    for cell in sheet["E"][1:]:
        cell.number_format = 'R$ #,##0.00'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in cells) + 2
        sheet.column_dimensions[get_column_letter(cells[0].column)].width = min(max(width, 12), 48)
    summary = workbook.create_sheet("Resumo")
    summary.append(("Indicador", "Quantidade"))
    summary.append(("OPERA", result.opera_count))
    summary.append(("Rede", result.rede_count))
    summary.append(("CMFlex", result.cmflex_count))
    summary.append(("Conciliados", result.matched_count))
    summary.append(("Divergentes", result.divergent_count))
    summary.append(("Rede ignorados", result.ignored_rede_count))
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
