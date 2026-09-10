from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def write_booking_csv(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)


def save_report_csv(
    path: Path, headers: list[str], records: list[dict[str, str]]
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(records)


def save_report_excel(
    path: Path, headers: list[str], records: list[dict[str, str]]
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conferência"
    sheet.append(headers)
    for record in records:
        sheet.append([record.get(header, "") for header in headers])
    fill = PatternFill("solid", fgColor="263238")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        content_width = max(len(str(cell.value or "")) for cell in cells) + 2
        sheet.column_dimensions[get_column_letter(cells[0].column)].width = min(
            max(content_width, 12), 48
        )
    workbook.save(path)
